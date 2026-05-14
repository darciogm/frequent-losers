"""
Gera dois .docx (Firpo e Soares) prontos para copiar e colar no Outlook.
Texto em parágrafos fluidos (sem quebras forçadas) — Outlook reformata
automaticamente. Calibri 11pt, espaçamento natural de email.
"""

from docx import Document
from docx.shared import Pt
from pathlib import Path

OUT = Path(__file__).parent

EMAILS = [
    {
        "filename": "email_firpo.docx",
        "to": "Sergio Firpo (Insper)",
        "subject": "Pedido de leitura presubmission — paper sobre set-asides de PME no JPubE",
        "salutacao": "Sergio,",
        "paragrafos": [
            "Boa tarde. Estou prestes a submeter ao Journal of Public Economics um paper solo que terminei de calibrar nas últimas semanas, e gostaria de pedir 30 minutos do seu tempo para um par de comentários antes da submissão — se couber na sua agenda nas próximas duas a três semanas.",
            "O paper usa a reversão da exceção do Grupo 65 na BEC-SP em março de 2018 como gatilho legal limpo para identificar o efeito da regra de exclusividade para ME/EPP em pregões eletrônicos. Combina a decomposição within-auction de Krasnokutskaya (2011) com um contrafactual de entrada endógena à la Athey-Coey-Levin (2013) e fecha com um cálculo de bem-estar no esquema de pesos de Saez-Stantcheva (2016). O resultado central é que ignorar a margem de entrada subestima o custo fiscal realizado em ~50-60%, e que a preferência de 10% domina o set-aside em mercados espessos mas não em mercados finos.",
            "O que eu queria mais era a sua leitura crítica de duas coisas específicas: (i) o framing da contribuição na introdução vis-à-vis a literatura de set-asides (Athey-Levin-Seira, Marion, Nakabayashi), e (ii) se a decomposição entre componente within-auction e margem de participação está clara o suficiente para um leitor de public economics que não vem de IO estrutural. Não preciso de revisão técnica linha a linha — só do filtro de quem já passou por esse pipeline editorial várias vezes.",
            "Anexei o PDF da v5 (45 páginas + apêndice online). Posso passar na sua sala em qualquer horário, ou a gente conversa por vídeo se for mais fácil.",
        ],
        "fechamento": ["Obrigado, e um abraço,", "Darcio"],
    },
    {
        "filename": "email_soares.docx",
        "to": "Rodrigo Soares (Insper)",
        "subject": "Pedido de leitura presubmission — paper para o JPubE, e talvez uma intro",
        "salutacao": "Rodrigo,",
        "paragrafos": [
            "Tudo bem? Estou prestes a submeter ao Journal of Public Economics um paper solo que terminei de calibrar. Antes de mandar, queria te pedir duas coisas — em ordem de importância.",
            "A primeira é a sua leitura crítica. O paper usa a reversão da exceção do Grupo 65 na BEC-SP em março de 2018 como gatilho legal limpo para identificar o efeito da regra de exclusividade para ME/EPP em pregões eletrônicos, decompõe o efeito entre a margem within-auction (Krasnokutskaya 2011) e a margem de participação endógena (Athey-Coey-Levin 2013), e fecha com cálculo de bem-estar no esquema Saez-Stantcheva. O resultado central é que ignorar a margem de entrada subestima o custo fiscal realizado em ~50-60%, e que a preferência de 10% domina o set-aside em mercados espessos mas não em finos. Não preciso de revisão linha a linha — só de 30 minutos do seu filtro sobre se a contribuição está bem posicionada para a literatura de public economics e se a identificação convence num primeiro round de leitura.",
            "A segunda — e essa só faz sentido se a primeira passar no seu filtro — é se você acharia razoável me apresentar ao Francesco Decarolis. Ele é o leitor Tier 2 mais natural para esse paper (procurement aplicado, mercados públicos europeus, recepção a working papers fora do circuito anglo), e você foi o caminho óbvio que me ocorreu. Se achar que ainda não está pronto, ou que faz mais sentido apresentar primeiro num seminário aqui, prefiro essa orientação à intro prematura.",
            "Anexei o PDF da v5 (45 páginas + apêndice). Posso passar na sua sala em qualquer momento da semana, ou a gente conversa por vídeo se for mais fácil.",
        ],
        "fechamento": ["Obrigado, e um abraço,", "Darcio"],
    },
]


def build_doc(spec: dict) -> Document:
    doc = Document()

    # estilo padrão Calibri 11
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # cabeçalho com Para / Assunto (para colar manualmente nos campos do Outlook)
    head_to = doc.add_paragraph()
    run = head_to.add_run("Para: ")
    run.bold = True
    head_to.add_run(spec["to"])

    head_subj = doc.add_paragraph()
    run = head_subj.add_run("Assunto: ")
    run.bold = True
    head_subj.add_run(spec["subject"])

    doc.add_paragraph("")  # linha em branco

    # corpo
    doc.add_paragraph(spec["salutacao"])
    doc.add_paragraph("")
    for par in spec["paragrafos"]:
        doc.add_paragraph(par)
        doc.add_paragraph("")
    for line in spec["fechamento"]:
        doc.add_paragraph(line)

    return doc


for spec in EMAILS:
    out_path = OUT / spec["filename"]
    build_doc(spec).save(out_path)
    print(f"escrito: {out_path}")
